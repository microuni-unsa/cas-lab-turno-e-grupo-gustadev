"""Completa Plan_Riesgo_L6.xlsx (solo celdas amarillas, no toca fórmulas azules)."""
import openpyxl

PATH = "l6/src/Plan_Riesgo_L6.xlsx"
wb = openpyxl.load_workbook(PATH)

# --- Activos: integrantes + propietario + C/I/D/A + justificación ---
ws = wb["Activos"]
ws["B5"] = "Sequeiros Condori Luis Gustavo"
activos = [
    # IDrow, propietario, C,I,D,A, justificación
    (10, "Secretaría Académica", 3, 3, 5, 3,
     "Servicio de matrícula: D5 por interrupción en periodo de matrícula; C3/I3/A3 moderados."),
    (11, "Responsable de datos", 5, 5, 4, 4,
     "Notas y datos personales: C5/I5 por divulgación o alteración masiva; D4 por bloqueo de consulta."),
    (12, "Desarrollo", 4, 4, 4, 4,
     "App web expuesta: C4/I4 por sesión y edición; D4 por caída en matrícula; A4 por suplantación."),
    (13, "Infraestructura", 3, 3, 5, 2,
     "Servidor y red: D5 por corte total del servicio; C3/I3 por acceso físico; A2 bajo."),
    (14, "Administrador TI", 5, 4, 4, 2,
     "Respaldos: C5 si se exfiltran; I4/D4 si no restauran; A2 bajo."),
    (15, "Seguridad TI", 4, 4, 3, 5,
     "Cuentas y personal: A5 por suplantación; C4/I4 por credenciales; D3 por bloqueo."),
]
for r, prop, c, i, d, a, just in activos:
    ws.cell(row=r, column=4, value=prop)
    ws.cell(row=r, column=6, value=c)
    ws.cell(row=r, column=7, value=i)
    ws.cell(row=r, column=8, value=d)
    ws.cell(row=r, column=9, value=a)
    ws.cell(row=r, column=10, value=just)

# --- Cruces: 8 escenarios R1-R8 (filas 6-13) ---
cruces = wb["Cruces Activo-Amenaza"]
escenarios = [
    ("R1", "A5", "T1", "Copias accesibles desde producción sin aislamiento", "D",
     "Ninguno", 4, 5,
     "P4: ransomware probable (>50% en 12m) sin copias aisladas; I5: >24h sin matrícula/notas y pérdida de respaldos."),
    ("R2", "A6", "T2", "Acceso sin segundo factor en cuentas académicas", "A",
     "Ninguno", 4, 4,
     "P4: phishing probable contra estudiantes/docentes; I4: 8-24h de suplantación y alteración relevante."),
    ("R3", "A2", "T3", "Permisos excesivos de edición de notas sin aprobación", "I",
     "Revisión manual esporádica", 4, 5,
     "P4: alteración probable por rol amplio; I5: alteración masiva de registros oficiales."),
    ("R4", "A4", "T4", "Sala sin UPS ni apagado controlado", "D",
     "Ninguno", 3, 4,
     "P3: corte posible (20-50%) en 12m; I4: 8-24h de servicio caído y apagado brusco."),
    ("R5", "A5", "T5", "Restauración nunca ensayada ni versionada", "D",
     "Versionado parcial", 3, 4,
     "P3: borrado accidental posible; I4: recuperación de 8-24h con registros por corregir."),
    ("R6", "A2", "T6", "Carpetas con enlace público sin restricción", "C",
     "Ninguno", 3, 5,
     "P3: divulgación posible por enlaces abiertos; I5: divulgación masiva de datos personales."),
    ("R7", "A4", "T7", "Servidor único sin redundancia ni plan de recuperación", "D",
     "Mantenimiento anual", 2, 4,
     "P2: falla poco probable (5-20%) con equipo nuevo; I4: 8-24h sin servicio si falla."),
    ("R8", "A3", "T10", "Despliegues sin prueba previa ni plan de reversión", "I",
     "Ninguno", 3, 3,
     "P3: error de actualización posible; I3: 4-8h degradado y corrección importante."),
]
for idx, (rid, aid, tid, vuln, dim, ctrl, p, imp, just) in enumerate(escenarios):
    r = 6 + idx
    cruces.cell(row=r, column=1, value=rid)
    cruces.cell(row=r, column=2, value=aid)
    cruces.cell(row=r, column=4, value=tid)
    cruces.cell(row=r, column=6, value=vuln)
    cruces.cell(row=r, column=7, value=dim)
    cruces.cell(row=r, column=8, value=ctrl)
    cruces.cell(row=r, column=9, value=p)
    cruces.cell(row=r, column=10, value=imp)
    cruces.cell(row=r, column=11, value=just)

# --- Análisis: tratamiento + residual (filas 31-38) ---
an = wb["Análisis de Riesgos"]
tratamientos = [
    ("Mitigar", "Copias aisladas 3-2-1 y ensayo de restauración", "Administrador TI", 30,
     "En ejecución", "Supuesto didáctico: copias aisladas en implantación, sin ensayo firmado.", 2, 3,
     "Nuevo tratamiento"),
    ("Mitigar", "Segundo factor y entrenamiento anti-phishing", "Seguridad TI", 15,
     "En ejecución", "Supuesto didáctico: 2FA en despliegue parcial, sin métrica de adopción.", 2, 3,
     "Nuevo tratamiento"),
    ("Mitigar", "Permisos por rol, aprobación y bitácora inmutable", "Desarrollo", 20,
     "Propuesto", "Supuesto didáctico: diseño de roles aprobado, pendiente de implementar.", 1, 5,
     "Nuevo tratamiento"),
    ("Mitigar", "UPS y procedimiento de apagado controlado", "Infraestructura", 30,
     "Propuesto", "Supuesto didáctico: UPS cotizado, pendiente de compra e instalación.", 2, 2,
     "Solicitar aceptación"),
    ("Mitigar", "Versionado y prueba de restauración trimestral", "Administrador de datos", 15,
     "Propuesto", "Supuesto didáctico: procedimiento redactado, sin prueba ejecutada.", 1, 3,
     "Solicitar aceptación"),
    ("Mitigar", "Acceso mínimo y revocación de enlaces públicos", "Responsable de datos", 15,
     "En ejecución", "Supuesto didáctico: revisión de permisos en curso, sin acta de cierre.", 1, 5,
     "Nuevo tratamiento"),
    ("Mitigar", "Servidor secundario y plan de recuperación", "Infraestructura", 45,
     "Propuesto", "Supuesto didáctico: arquitectura propuesta, sin compra ni pruebas.", 1, 4,
     "Solicitar aceptación"),
    ("Mitigar", "Pipeline de pruebas y plan de reversión por despliegue", "Desarrollo", 20,
     "Propuesto", "Supuesto didáctico: pipeline diseñado, sin ejecución registrada.", 1, 3,
     "Solicitar aceptación"),
]
for idx, (trat, acc, resp, plazo, estado, evid, pr, ir, dec) in enumerate(tratamientos):
    r = 31 + idx
    an.cell(row=r, column=2, value=trat)
    an.cell(row=r, column=3, value=acc)
    an.cell(row=r, column=4, value=resp)
    an.cell(row=r, column=5, value=plazo)
    an.cell(row=r, column=6, value=estado)
    an.cell(row=r, column=7, value=evid)
    an.cell(row=r, column=8, value=pr)
    an.cell(row=r, column=9, value=ir)
    an.cell(row=r, column=14, value=dec)

wb.save(PATH)
print("OK guardado", PATH)
